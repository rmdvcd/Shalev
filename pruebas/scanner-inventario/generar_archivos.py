#!/usr/bin/env python3
"""
Genera los 5 archivos de prueba manual del escáner de inventario de Shalev
(Inventario -> «Escanear o importar inventario»).

    python3 generar_archivos.py

Requiere: pillow, reportlab, openpyxl, python-docx.
Las tipografías manuscritas (OFL) se buscan en ./fuentes, /tmp/fuentes o /usr/share/fonts;
si no están, se descargan del repositorio google/fonts vía api.github.com.
Los archivos se escriben junto a este script.
"""

from __future__ import annotations

import base64
import json
import math
import os
import random
import urllib.request
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

AQUI = Path(__file__).resolve().parent
FUENTES_LOCALES = [
    AQUI / "fuentes",
    Path("/tmp/fuentes"),
    Path("/usr/share/fonts/truetype/dejavu"),
]

FONT_GOOGLE = {
    "PatrickHand-Regular.ttf": "ofl/patrickhand/PatrickHand-Regular.ttf",
    "ArchitectsDaughter-Regular.ttf": "ofl/architectsdaughter/ArchitectsDaughter-Regular.ttf",
}


# --------------------------------------------------------------------------- #
# Utilidades
# --------------------------------------------------------------------------- #
def ruta_fuente(nombre: str) -> Path:
    for carpeta in FUENTES_LOCALES:
        candidato = carpeta / nombre
        if candidato.exists():
            return candidato
    destino = Path("/tmp/fuentes") / nombre
    if destino.exists():
        return destino
    destino.parent.mkdir(parents=True, exist_ok=True)
    url = f"https://api.github.com/repos/google/fonts/contents/{FONT_GOOGLE[nombre]}"
    with urllib.request.urlopen(url, timeout=60) as r:
        datos = json.load(r)
    destino.write_bytes(base64.b64decode(datos["content"]))
    print(f"  · tipografía descargada: {nombre}")
    return destino


def fuente_dejavu(nombre: str) -> Path:
    for carpeta in FUENTES_LOCALES:
        candidato = carpeta / nombre
        if candidato.exists():
            return candidato
    raise FileNotFoundError(nombre)


# --------------------------------------------------------------------------- #
# 1. Manuscrito fotografiado (JPG)
# --------------------------------------------------------------------------- #
def manuscrito(destino: Path) -> None:
    rng = random.Random(28)
    W, H = 1600, 1800
    papel = Image.new("RGB", (W, H), (250, 248, 239))
    dibujo = ImageDraw.Draw(papel)

    # Degradado de luz (foto con flash suave desde arriba a la izquierda).
    for y in range(H):
        t = y / H
        tono = int(6 * t)
        dibujo.line([(0, y), (W, y)], fill=(250 - tono, 248 - tono, 239 - tono))

    # Renglones de libreta + margen rojo.
    for y in range(150, H - 40, 96):
        dibujo.line([(0, y), (W, y)], fill=(206, 219, 235), width=2)
    dibujo.line([(132, 0), (132, H)], fill=(214, 156, 156), width=3)

    tinta = (28, 40, 82)
    tinta_suave = (38, 52, 96)
    f_titulo = ImageFont.truetype(str(ruta_fuente("ArchitectsDaughter-Regular.ttf")), 62)
    f_sub = ImageFont.truetype(str(ruta_fuente("PatrickHand-Regular.ttf")), 40)
    f_cab = ImageFont.truetype(str(ruta_fuente("PatrickHand-Regular.ttf")), 44)
    f_dato = ImageFont.truetype(str(ruta_fuente("PatrickHand-Regular.ttf")), 46)

    def escribir(palabras, x, y, tam, color, sombra=False):
        """Escribe una lista de palabras con temblor, rotación e inclinación leves."""
        cursor = x
        for palabra in palabras:
            f = ImageFont.truetype(str(ruta_fuente("PatrickHand-Regular.ttf")), tam + rng.randint(-2, 2))
            caja = f.getbbox(palabra)
            ancho = caja[2] - caja[0]
            alto = caja[3] - caja[1]
            capa = Image.new("RGBA", (ancho + 26, alto + 26), (0, 0, 0, 0))
            # Se compensa el recuadro real del glifo: si no, letras sin ascendente
            # (la «u» de unidad) se recortan contra el borde inferior.
            ImageDraw.Draw(capa).text((13 - caja[0], 13 - caja[1]), palabra, font=f, fill=color + (255,))
            if sombra:
                capa = capa.filter(ImageFilter.GaussianBlur(0.7))
            capa = capa.rotate(rng.uniform(-2.4, 2.4), resample=Image.BICUBIC, expand=True)
            dy = y + rng.randint(-3, 3)
            papel.paste(capa, (int(cursor), int(dy)), capa)
            cursor += ancho + rng.randint(7, 13)

    ancho_texto = lambda texto, tam: ImageFont.truetype(
        str(ruta_fuente("PatrickHand-Regular.ttf")), tam
    ).getlength(texto)

    # Encabezado escrito a mano.
    escribir(["INVENTARIO", "BODEGA"], 168, 58, 62, tinta, sombra=True)
    escribir(["28", "oct", "2026"], 900, 66, 40, tinta_suave)

    # Cabecera de columnas (subrayada a mano).
    columnas = [
        (168, ["Producto"]),
        (820, ["Cant"]),
        (975, ["Unidad"]),
        (1160, ["P", "venta"]),
        (1370, ["Costo"]),
    ]
    for x, palabras in columnas:
        escribir(palabras, x, 206, 42, tinta)
    dibujo.line([(168, 290), (1452, 290)], fill=tinta, width=3)
    dibujo.line([(168, 294), (1290, 294)], fill=tinta, width=2)

    filas = [
        (["Café", "Serrano"], "30", "", "850", "600"),
        (["Refresco", "Cola"], "24", "", "250", "180"),
        (["Pan", "de", "molde"], "12", "", "320", "250"),
        (["Queso", "Gouda"], "6", "", "950", "700"),
        (["Aceite", "girasol"], "10", "", "1200", "900"),
        (["Jabón", "de", "baño"], "15", "", "150", "95"),
        (["Leche", "polvo"], "8", "", "1650", "1300"),
        (["Yogurt", "natural"], "20", "", "180", "120"),
        (["Azúcar"], "2,5", "kg", "", "350"),
        (["Harina"], "5", "kg", "", "480"),
        (["Café", "en", "grano"], "1", "kg", "", "2400"),
    ]

    y = 330
    for i, (nombre, cant, unidad, pv, pc) in enumerate(filas):
        if i == 8:  # separación escrita a mano antes de los insumos
            y += 16
            escribir(["INSUMOS"], 168, y, 40, tinta)
            dibujo.line([(168, y + 58), (520, y + 58)], fill=tinta, width=2)
            y += 62
        escribir(nombre, 168, y, 48, tinta)
        escribir([cant], 828, y, 48, tinta)
        # Todas las filas llevan las cinco columnas escritas: así la vista
        # «Ver como tabla» mantiene alineadas las columnas.
        escribir([unidad if unidad else "u"], 984, y, 48, tinta)
        escribir([pv if pv else "—"], 1140, y, 48, tinta)
        escribir([pc], 1360, y, 48, tinta)
        y += 96

    # Nota suelta al pie: texto que el usuario debe aprender a descartar.
    escribir(["revisar", "precios", "del", "queso"], 190, y + 36, 40, (150, 40, 40))
    escribir(["falta", "leche", "en", "polvo"], 190, y + 88, 40, (150, 40, 40))

    # Textura de papel + grano de la cámara.
    grano = Image.effect_noise((W // 3, H // 3), 26).resize((W, H)).convert("RGB")
    papel = Image.blend(papel, grano, 0.055)
    # Sombra en los bordes (foto tomada con el teléfono encima).
    vineta = Image.new("L", (W, H), 0)
    ImageDraw.Draw(vineta).ellipse([-W * 0.22, -H * 0.18, W * 1.22, H * 1.18], fill=255)
    vineta = vineta.filter(ImageFilter.GaussianBlur(180))
    oscuro = Image.new("RGB", (W, H), (196, 190, 176))
    papel = Image.composite(papel, oscuro, vineta)
    papel = papel.rotate(-1.6, resample=Image.BICUBIC, expand=True, fillcolor=(238, 234, 224))
    papel = papel.filter(ImageFilter.GaussianBlur(0.5))
    papel.save(destino, "JPEG", quality=90, optimize=True)
    print(f"  · {destino.name} ({destino.stat().st_size // 1024} KB)")


# --------------------------------------------------------------------------- #
# 2. PDF impreso de 2 páginas
# --------------------------------------------------------------------------- #
PAGINA_1 = [
    ("Martillo de acero", "Herramientas", "8", "u", "1 250,00", "890,00"),
    ("Cinta métrica 5 m", "Herramientas", "15", "u", "650,00", "420,00"),
    ("Tornillos 2 pulg", "Ferretería", "40", "u", "180,00", "110,00"),
    ("Pintura blanca 4 L", "Pintura", "12", "u", "2 100,00", "1 550,00"),
    ("Brocha 3 pulg", "Pintura", "20", "u", "320,00", "190,00"),
    ("Llave de agua 1/2", "Plomería", "25", "u", "480,00", "300,00"),
    ("Cemento gris", "Construcción", "60", "u", "95,00", "62,00"),
    ("Alambre dulce", "Eléctrico", "9", "u", "780,00", "510,00"),
]

PAGINA_2 = [
    ("Thinner acrílico", "Insumos", "5", "L", "—", "1 100,00"),
    ("Aceite de linaza", "Insumos", "3,5", "L", "—", "1 480,00"),
    ("Yeso fino", "Insumos", "12", "kg", "—", "210,00"),
    ("Sellador madera", "Insumos", "6,25", "L", "—", "890,00"),
    ("Cemento blanco", "Insumos", "20", "kg", "—", "265,00"),
    ("Resina epóxica", "Insumos", "1,5", "kg", "—", "3 400,00"),
]


def pdf(destino: Path) -> None:
    pdfmetrics.registerFont(TTFont("DejaVu", str(fuente_dejavu("DejaVuSans.ttf"))))
    pdfmetrics.registerFont(TTFont("DejaVuBold", str(fuente_dejavu("DejaVuSans-Bold.ttf"))))
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=estilos["Title"], fontName="DejaVuBold", fontSize=15, spaceAfter=2)
    sub = ParagraphStyle("s", parent=estilos["Normal"], fontName="DejaVu", fontSize=9.5,
                         textColor=colors.HexColor("#444444"), spaceAfter=8)
    nota = ParagraphStyle("n", parent=estilos["Normal"], fontName="DejaVu", fontSize=8,
                          textColor=colors.HexColor("#555555"), spaceBefore=6)

    estilo_cabecera = ParagraphStyle(
        "cab", parent=estilos["Normal"], fontName="DejaVuBold", fontSize=8,
        leading=10, alignment=1, textColor=colors.white,
    )

    def tabla(filas):
        datos = [[Paragraph(t, estilo_cabecera) for t in (
            "Producto", "Categoría", "Existencia", "Unidad",
            "Precio venta (CUP)", "Precio costo (CUP)")]]
        datos += [list(f) for f in filas]
        t = Table(datos, colWidths=[40 * mm, 24 * mm, 25 * mm, 17 * mm, 35 * mm, 35 * mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("FONTNAME", (0, 1), (-1, -1), "DejaVu"),
            ("FONTSIZE", (0, 1), (-1, -1), 11),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#455A64")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4F5")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA5AB")),
            ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 11),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 11),
        ]))
        return t

    doc = SimpleDocTemplate(
        str(destino), pagesize=A4,
        leftMargin=16 * mm, rightMargin=16 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
        title="Inventario impreso de prueba", author="Shalev",
    )
    historia = [
        Paragraph("Ferretería Los Pinos", titulo),
        Paragraph("Inventario físico — 28 de octubre de 2026 — Página 1 de 2", sub),
        tabla(PAGINA_1),
        Spacer(1, 8),
        Paragraph("Los precios están en CUP. Las existencias se contaron en el almacén.", nota),
        PageBreak(),
        Paragraph("Ferretería Los Pinos", titulo),
        Paragraph("Insumos de taller — 28 de octubre de 2026 — Página 2 de 2", sub),
        tabla(PAGINA_2),
        Spacer(1, 8),
        Paragraph("Insumos: la cantidad se lleva en kg o L, no en unidades enteras.", nota),
    ]
    doc.build(historia)
    print(f"  · {destino.name} ({destino.stat().st_size // 1024} KB)")


# --------------------------------------------------------------------------- #
# 3. Libro Excel (.xlsx) con dos hojas
# --------------------------------------------------------------------------- #
def xlsx(destino: Path) -> None:
    wb = Workbook()
    cabecera = PatternFill("solid", fgColor="E8EDF0")
    borde = Border(*[Side(style="thin", color="B0BEC5")] * 4)
    negrita = Font(bold=True, color="1F2A30")

    hoja = wb.active
    hoja.title = "Inventario"
    hoja.append(["Producto", "Categoría", "Existencia", "Unidad", "Precio venta", "Precio costo", "Descripción"])
    datos = [
        ("Café Serrano", "Alimentos", 30, "u", 850, 600, "Paquete 250 g"),
        ("Refresco Cola", "Bebidas", 24, "u", 250, 180, "Lata 355 ml"),
        ("Pan de molde", "Alimentos", 12, "u", 320, 250, ""),
        ("Queso Gouda", "Alimentos", 6, "u", 950, 700, ""),
        ("Detergente", "Limpieza", 18, "u", 275.5, 160.25, "Bolsa 500 g"),
        ("Yogurt natural", "Alimentos", 20, "u", 180, 120, ""),
        ("Pizza familiar", "Elaborado", 0, "u", 1500, None, "Se vende por receta"),
        ("Jabón de baño", "Limpieza", 15, "u", 150, 95, ""),
        ("Sin precio", "Bebidas", 10, "u", None, 100, "Fila que debe omitirse"),
        ("Lata de atún", "Alimentos", 0, "u", 410, 320, "Sin existencias"),
        ("Chocolate", "Dulces", 500, "u", 95, 60,
         "Barra de 100 g con descripción larga para probar el recorte a 40 caracteres"),
    ]
    for fila in datos:
        hoja.append(list(fila))

    hoja2 = wb.create_sheet("Insumos")
    hoja2.append(["Nombre", "Unidad", "Cantidad", "Costo", "Precio venta"])
    for fila in [
        ("Azúcar", "kg", 2.5, 350, None),
        ("Harina", "kg", 5, 480, None),
        ("Aceite vegetal", "L", 8.75, 620, 900),
        ("Café en grano", "kg", 1, 2400, None),
        ("Sal", "kg", 3, 120, 150),
    ]:
        hoja2.append(list(fila))

    for h in wb.worksheets:
        ancho = {"A": 26, "B": 16, "C": 12, "D": 9, "E": 14, "F": 14, "G": 46}
        for col in h[1]:
            col.font = negrita
            col.fill = cabecera
            col.border = borde
            col.alignment = Alignment(horizontal="center")
        for fila in h.iter_rows(min_row=2):
            for celda in fila:
                celda.border = borde
        for letra, valor in ancho.items():
            h.column_dimensions[letra].width = valor
        for fila in h.iter_rows(min_row=2, min_col=5, max_col=6):
            for celda in fila:
                if isinstance(celda.value, (int, float)):
                    celda.number_format = "#,##0.00"
        if h.title == "Inventario":
            for fila in h.iter_rows(min_row=2, min_col=3, max_col=3):
                fila[0].number_format = "0"

    wb.save(destino)
    print(f"  · {destino.name} ({destino.stat().st_size // 1024} KB)")


# --------------------------------------------------------------------------- #
# 4. CSV guardado por Excel (BOM + punto y coma + comas decimales)
# --------------------------------------------------------------------------- #
CSV_FILAS = [
    'Miel de abeja;12;u;450,00;310,00;Alimentos;"Frasco 250 ml; tapa roja"',
    'Café molido;25;u;750,00;520,00;Alimentos;Paquete 500 g',
    'Aceite de oliva;3;L;1 250,00;980,00;Insumos;"Botella; 1 L"',
    'Harina de trigo;10;kg;95,00;70,00;Insumos;',
    'Té verde;40;u;120,00;80,00;Bebidas;',
    'Chocolate en polvo;8;u;640,00;450,00;Alimentos;Lata 400 g',
    'Vinagre;0;L;300,00;210,00;Insumos;Sin existencias',
    'Galletas surtidas;36;u;75,00;50,00;Dulces;"Sobre individual; surtido"',
]


def csv(destino: Path) -> None:
    # Igual que «Guardar como CSV» de Excel: BOM UTF-8 y saltos CRLF.
    contenido = "Artículo;Existencia;Unidad;PVP;Costo;Categoría;Notas\r\n" + "\r\n".join(CSV_FILAS) + "\r\n"
    destino.write_bytes(b"\xef\xbb\xbf" + contenido.encode("utf-8"))
    print(f"  · {destino.name} ({destino.stat().st_size // 1024} KB)")


# --------------------------------------------------------------------------- #
# 5. Documento Word (.docx) con dos tablas
# --------------------------------------------------------------------------- #
def docx(destino: Path) -> None:
    doc = Document()
    doc.add_heading("Panadería La Espiga", level=1)
    p = doc.add_paragraph("Inventario del 28/10/2026 — conteo de la tarde")
    p.runs[0].font.size = Pt(10)

    tabla1 = doc.add_table(rows=1, cols=5)
    tabla1.style = "Light Grid Accent 1"
    for celda, texto in zip(tabla1.rows[0].cells,
                            ["Producto", "Categoría", "Cantidad", "Precio venta", "Precio costo"]):
        celda.text = texto
    filas1 = [
        ("Pan de agua", "Alimentos", "40", "25,00", "15,00"),
        ("Pan dulce", "Alimentos", "30", "35,00", "22,00"),
        ("Croissant", "Alimentos", "12", "60,00", "40,00"),
        ("Pastel de chocolate", "Dulces", "2", "850,00", "600,00"),
        ("", "", "", "", ""),
        ("Café molido", "Bebidas", "10", "750,00", "520,00"),
        ("Refresco", "Bebidas", "24", "250,00", "180,00"),
        ("Galletas", "Dulces", "18", "90,00", "60,00"),
        ("Flan", "Dulces", "6", "", "300,00"),
    ]
    for fila in filas1:
        celdas = tabla1.add_row().cells
        for celda, valor in zip(celdas, fila):
            celda.text = valor

    doc.add_paragraph("Insumos de la panadería")
    tabla2 = doc.add_table(rows=1, cols=4)
    tabla2.style = "Light Grid Accent 1"
    for celda, texto in zip(tabla2.rows[0].cells, ["Nombre", "Unidad", "Cantidad", "Costo"]):
        celda.text = texto
    for fila in [
        ("Harina", "kg", "25", "480,00"),
        ("Azúcar", "kg", "12,5", "350,00"),
        ("Levadura", "g", "500", "8,00"),
        ("Mantequilla", "kg", "6", "1 250,00"),
    ]:
        celdas = tabla2.add_row().cells
        for celda, valor in zip(celdas, fila):
            celda.text = valor

    doc.add_paragraph("Revisar el flan: vence el 30/10 y no tiene precio de venta.")
    doc.save(destino)
    print(f"  · {destino.name} ({destino.stat().st_size // 1024} KB)")


def main() -> None:
    print("Generando archivos de prueba del escáner de inventario…")
    manuscrito(AQUI / "01-inventario-manuscrito.jpg")
    pdf(AQUI / "02-inventario-impreso-2-paginas.pdf")
    xlsx(AQUI / "03-inventario-excel.xlsx")
    csv(AQUI / "04-inventario.csv")
    docx(AQUI / "05-inventario-word.docx")
    print("Listo.")


if __name__ == "__main__":
    main()
